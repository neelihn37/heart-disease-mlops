from docx import Document

doc = Document()
doc.add_heading('End-to-End Heart Disease MLOps Project Report', 0)

sections = [
    ("1. Introduction & Executive Summary", "Overview of the heart disease classification MLOps system, business value, and objectives."),
    ("2. Architecture & Tech Stack", "Detailed breakdown of Minikube, Kubernetes, Docker, FastAPI, MLflow, and GitHub Actions."),
    ("3. Data Pipeline & Exploratory Data Analysis (EDA)", "Dataset characteristics, feature distributions, missing values handling, and correlations."),
    ("4. Model Training & Experiment Tracking", "Scikit-learn model selection, hyperparameter tuning, evaluation metrics (Accuracy, ROC-AUC), and MLflow tracking."),
    ("5. Containerization & API Development", "FastAPI implementation, Pydantic schema validation, Docker multi-stage builds, and container execution."),
    ("6. Kubernetes Orchestration & Deployment", "Minikube deployment manifests, NodePort configurations, pod health checks, and service routing."),
    ("7. CI/CD Automation & Testing", "GitHub Actions workflow setup, automated unit testing with pytest, and regression checks."),
    ("8. Monitoring, Logging & Maintenance", "Drift detection strategy, model retraining triggers, and cluster logging setup."),
    ("9. Results & Discussion", "Performance benchmarks, latency evaluations, and API verification results."),
    ("10. Conclusion & Future Scope", "Project summary, scaling considerations to enterprise cloud clusters (EKS/GKE), and future enhancements.")
]

for title, desc in sections:
    doc.add_heading(title, level=1)
    doc.add_paragraph(desc)
    doc.add_paragraph("[ Insert relevant screenshots, metrics charts, or code snippets here ]\n")

doc.save("Report_Heart_Disease_MLOps.docx")
print("Generated Report_Heart_Disease_MLOps.docx successfully!")